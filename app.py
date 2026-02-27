import streamlit as st
import pandas as pd
from docxtpl import DocxTemplate
import jinja2
from docx import Document
from docxcompose.composer import Composer
import os
import shutil
from pathlib import Path
from datetime import datetime
import tempfile
import zipfile
import traceback

# Page config must be the first Streamlit command
st.set_page_config(
    page_title="Document Generator",
    page_icon="📄",
    layout="wide",
    initial_sidebar_state="expanded"
)

def replace_text(text):
    """Replace _x000D_ with . in text"""
    if isinstance(text, str):
        return text.replace("_x000D_", ".")
    return text

def process_dataframe(df):
    """Process dataframe to clean text"""
    records = df.to_dict(orient='records')
    for record in records:
        for key, value in record.items():
            record[key] = replace_text(value)
    return records

def generate_documents(excel_file, template_file, output_filename=None, 
                       keep_individual=False, create_backup=True):
    """
    Generate documents from Excel and merge them
    """
    try:
        # Read Excel data
        df = pd.read_excel(excel_file)
        records = process_dataframe(df)
        
        # Create progress elements
        progress_bar = st.progress(0)
        status_text = st.empty()
        
        generated_files = []
        temp_dir = Path("temp_individual_docs")
        
        if not keep_individual:
            temp_dir.mkdir(exist_ok=True)
        
        # Generate individual documents
        for i, context in enumerate(records):
            status_text.text(f"Processing record {i+1} of {len(records)}")
            
            doc = DocxTemplate(template_file)
            jinja_env = jinja2.Environment(autoescape=True)
            doc.render(context, jinja_env)
            
            if keep_individual:
                out_name = f"generated_doc_{i+1:03d}.docx"
            else:
                out_name = temp_dir / f"doc_{i+1:03d}.docx"
            
            doc.save(str(out_name))
            generated_files.append(str(out_name))
            progress_bar.progress((i + 1) / len(records))
        
        # Set output filename
        if output_filename is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_filename = f"merged_document_{timestamp}.docx"
        
        # Create backup if exists
        if create_backup and os.path.exists(output_filename):
            backup_name = f"backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{output_filename}"
            shutil.copy2(output_filename, backup_name)
            st.info(f"✅ Backup created: {backup_name}")
        
        # Merge documents
        status_text.text("Merging documents...")
        master = Document(generated_files[0])
        composer = Composer(master)
        
        for file in generated_files[1:]:
            doc = Document(file)
            composer.append(doc)
        
        composer.save(output_filename)
        
        # Cleanup
        if not keep_individual:
            for file in generated_files:
                try:
                    os.remove(file)
                except:
                    pass
            try:
                if temp_dir.exists() and not any(temp_dir.iterdir()):
                    temp_dir.rmdir()
            except:
                pass
        
        status_text.text("✅ Complete!")
        progress_bar.empty()
        
        return output_filename, generated_files if keep_individual else None
        
    except Exception as e:
        st.error(f"Error during generation: {str(e)}")
        st.code(traceback.format_exc())
        return None, None

def main():
    st.title("📄 Document Generator & Merger")
    st.markdown("---")
    
    # Sidebar for instructions
    with st.sidebar:
        st.header("ℹ️ Instructions")
        st.markdown("""
        1. **Upload Excel file** containing your data
        2. **Upload Word template** with Jinja2 placeholders
        3. Configure options
        4. Click **Generate Documents**
        5. Download your merged document
        
        **Excel format:**
        - First row should be column headers
        - Each row becomes one document
        
        **Template placeholders:**
        - Use `{{ column_name }}` in your Word template
        - Example: `{{ Name }}`, `{{ Date }}`
        """)
        
        st.header("⚙️ Options")
        keep_individual = st.checkbox(
            "Keep individual files", 
            value=False,
            help="Save each generated document separately"
        )
        create_backup = st.checkbox(
            "Create backup if file exists", 
            value=True,
            help="Create a backup if the output file already exists"
        )
        custom_filename = st.text_input(
            "Custom output filename (optional)",
            placeholder="my_document.docx",
            help="Leave empty for auto-generated filename"
        )
    
    # Main content area
    col1, col2 = st.columns(2)
    
    with col1:
        st.subheader("📁 Upload Files")
        
        # File uploaders
        excel_file = st.file_uploader(
            "Choose Excel file",
            type=['xlsx', 'xls'],
            help="Upload your Excel file with data"
        )
        
        template_file = st.file_uploader(
            "Choose Word template",
            type=['docx'],
            help="Upload your Word template with Jinja2 placeholders"
        )
        
        # Show file info
        if excel_file:
            st.success(f"✅ Excel: {excel_file.name}")
            # Preview Excel data
            try:
                df = pd.read_excel(excel_file)
                st.caption(f"Found {len(df)} records with columns: {', '.join(df.columns)}")
                
                # Show sample data
                with st.expander("Preview Excel Data (first 3 rows)"):
                    st.dataframe(df.head(3))
            except Exception as e:
                st.error(f"Error reading Excel: {e}")
        
        if template_file:
            st.success(f"✅ Template: {template_file.name}")
    
    with col2:
        st.subheader("🚀 Generate")
        
        # Process button
        if excel_file and template_file:
            if st.button("Generate Documents", type="primary", use_container_width=True):
                try:
                    # Save uploaded files temporarily
                    with tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx') as tmp_excel:
                        tmp_excel.write(excel_file.getvalue())
                        excel_path = tmp_excel.name
                    
                    with tempfile.NamedTemporaryFile(delete=False, suffix='.docx') as tmp_template:
                        tmp_template.write(template_file.getvalue())
                        template_path = tmp_template.name
                    
                    # Process
                    with st.spinner("Generating documents..."):
                        output_file, individual_files = generate_documents(
                            excel_file=excel_path,
                            template_file=template_path,
                            output_filename=custom_filename if custom_filename else None,
                            keep_individual=keep_individual,
                            create_backup=create_backup
                        )
                    
                    if output_file:
                        st.success("✅ Documents generated successfully!")
                        
                        # Download section
                        st.subheader("📥 Download Results")
                        
                        # Download merged file
                        with open(output_file, 'rb') as f:
                            st.download_button(
                                label="📄 Download Merged Document",
                                data=f,
                                file_name=os.path.basename(output_file),
                                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                use_container_width=True
                            )
                        
                        # If keeping individual files, create a zip
                        if keep_individual and individual_files:
                            zip_path = "individual_documents.zip"
                            with zipfile.ZipFile(zip_path, 'w') as zipf:
                                for file in individual_files:
                                    zipf.write(file, os.path.basename(file))
                            
                            with open(zip_path, 'rb') as f:
                                st.download_button(
                                    label="📦 Download Individual Files (ZIP)",
                                    data=f,
                                    file_name="individual_documents.zip",
                                    mime="application/zip",
                                    use_container_width=True
                                )
                            
                            # Cleanup individual files
                            for file in individual_files:
                                os.remove(file)
                            os.remove(zip_path)
                        
                        # Cleanup temp files
                        os.remove(excel_path)
                        os.remove(template_path)
                        if os.path.exists(output_file):
                            os.remove(output_file)  # Clean up after download
                            
                except Exception as e:
                    st.error(f"Error: {str(e)}")
                    st.code(traceback.format_exc())
        else:
            st.info("👆 Please upload both Excel and template files to begin")
    
    # Footer
    st.markdown("---")
    st.markdown("Made with ❤️ using Streamlit")

if __name__ == "__main__":
    main()